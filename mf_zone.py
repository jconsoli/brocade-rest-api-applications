#!/usr/bin/python
# -*- coding: utf-8 -*-
"""
Copyright 2024, 2025, 2026 Jack Consoli.  All rights reserved.

**License**

Licensed under the Apache License, Version 2.0 (the "License"); you may not use this file except in compliance with
the License. You may also obtain a copy of the License at https://www.apache.org/licenses/LICENSE-2.0

Unless required by applicable law or agreed to in writing, software distributed under the License is distributed on an
"AS IS" BASIS, WITHOUT WARRANTIES OR CONDITIONS OF ANY KIND, either express or implied. See the License for the specific
language governing permissions and limitations under the License.

The license is free for single customer use (internal applications). Use of this module in the production,
redistribution, or service delivery for commerce requires an additional license. Contact jack_consoli@yahoo.com for
details.

**Description**

Creates a report in Excel Workbook format from a brcddb project

**Version Control**

+-----------+---------------+---------------------------------------------------------------------------------------+
| Version   | Last Edit     | Description                                                                           |
+===========+===============+=======================================================================================+
| 1.0.0     | 26 Dec 2024   | Initial Launch                                                                        |
+-----------+---------------+---------------------------------------------------------------------------------------+
| 1.0.1     | 01 Mar 2025   | Help and error message enhancements.                                                  |
+-----------+---------------+---------------------------------------------------------------------------------------+
| 1.0.2     | 25 Aug 2025   | Use brcddb.util.util.get_import_modules to dynamically determined imported libraries. |
+-----------+---------------+---------------------------------------------------------------------------------------+
| 1.0.3     | 19 Oct 2025   | Updated comments only.                                                                |
+-----------+---------------+---------------------------------------------------------------------------------------+
| 1.0.4     | 20 Feb 2026   | Updated copyright notice.                                                             |
+-----------+---------------+---------------------------------------------------------------------------------------+
| 1.0.5     | 10 Mar 2026   | Consolidated project read and Excel error messaging.                                  |
+-----------+---------------+---------------------------------------------------------------------------------------+
| 1.0.6     | 01 Aug 2026   | Modified to use new brcddb_switch.best_switch_name() format.                          |
+-----------+---------------+---------------------------------------------------------------------------------------+
| 1.0.7     | 26 Oct 2026   | Moved some private functions from here to brcddp.brcddb_fabric.py and made public.    |
+-----------+---------------+---------------------------------------------------------------------------------------+
"""
__author__ = 'Jack Consoli'
__copyright__ = 'Copyright 2024, 2025, 2026 Jack Consoli'
__date__ = '26 Oct 2026'
__license__ = 'Apache License, Version 2.0'
__email__ = 'jack_consoli@yahoo.com'
__maintainer__ = 'Jack Consoli'
__status__ = 'Released'
__version__ = '1.0.7'

import os
import openpyxl.utils.cell as xl
import brcdapi.log as brcdapi_log
import brcdapi.fos_auth as fos_auth
import brcdapi.file as brcdapi_file
import brcdapi.excel_fonts as excel_fonts
import brcdapi.excel_util as excel_util
import brcdapi.gen_util as gen_util
import brcdapi.util as brcdapi_util
import brcddb.brcddb_project as brcddb_project
import brcddb.brcddb_fabric as brcddb_fabric
import brcddb.brcddb_common as brcddb_common
import brcddb.brcddb_zone as brcddb_zone
import brcddb.util.iocp as brcddb_iocp
import brcddb.report.utils as report_utils

_DOC_STRING = False  # Should always be False. Prohibits any code execution. Only useful for building documentation
# _STAND_ALONE: True: Executes as a standalone module taking input from the command line. False: Does not automatically
# execute. This is useful when importing this module into another module that calls psuedo_main().
_STAND_ALONE = True  # See note above

_std_font = excel_fonts.font_type('std')
_bold_font = excel_fonts.font_type('bold')
_align_wrap = excel_fonts.align_type('wrap')
_zonecfg_col = (20, 68)


# For input parameter definitions, search for _input_d
class Found (Exception):
    pass


def _zone_option_all_func(fabric_obj, zonecfg, rnid_d, notes_l):
    """Returns a list with a single zone dictionary that includes all ports except ICL ports

    :param fabric_obj: Fabric object
    :type fabric_obj: brcddb.classes.fabric.FabricObj
    :param zonecfg: Zone configuration name
    :type zonecfg: str
    :param rnid_d: RNID dictionary. Output of brcddb_fabric.rnid_groups()
    :type rnid_d: dict
    :param notes_l: List of notes to add to the notes page
    :type notes_l: list
    :return: Dictionaries with zone DB definitions for use with brcddb.brcddb_zone.add_zone_worksheet()
    :rtype: list
    """

    return brcddb_fabric.ficon_zone_all(fabric_obj, zonecfg, notes_l)


def _zone_option_all_f_func(fabric_obj, zonecfg, rnid_d, notes_l):
    """Returns a zone dictionary with a single zone that includes all F ports. See _zone_option_all_func"""

    return brcddb_fabric.ficon_zone_all_f(fabric_obj, zonecfg, notes_l)


def _zone_option_dev_func(fabric_obj, zonecfg, rnid_d, notes_l):
    """Returns a list of zones for each CHPID with access to all control units. See _zone_option_all_func"""

    return brcddb_fabric.ficon_zone_by_device(fabric_obj, zonecfg, rnid_d, notes_l)


def _zone_option_channel_func(fabric_obj, zonecfg, rnid_d, notes_l):
    """Returns list of dict specific for each channel and associated link addresses. See _zone_option_all_func"""

    return brcddb_fabric.ficon_zone_by_channel(fabric_obj, zonecfg, notes_l)


# _zone_options_l: List of dictionaries as follows:
# h     Help message used with _eh_l
# a     Pointer to the function that builds the report
# sname Sheet name
_zone_options_l = (
    dict(
        h=dict(p1='A single large standard zone with all physical ports in the fabric except ICL ports is created. '
                  'This is the easiest zoning method to use because it will support any IODF. It only needs to be '
                  'created and activated once.'),
        a=_zone_option_all_func,
        sname='All_Ports',
    ),
    dict(
        h=dict(p1='The same as "All_Ports" except only online ports that are not E-Ports or ports configured for test '
                  'or loopback are included in the zone. Any port associated with a channel or control unit interface '
                  'that was offline when the data was collected will need to be added.'),
        a=_zone_option_all_f_func,
        sname='All_F_Ports',
    ),
    dict(
        h=[
            dict(p1='Device zoning creates a group of control units by serial number. A separate peer zone for each '
                    'CHPID and each group of control units is created.'),
            dict(p1=''),
            dict(p1='**WARNING** A separate zone for Channel-to-Channel (CTC) will need to be created. Similarly, '
                    'peer-to-peer disk mirroring ports will need a separate zone as well. Since peer-to-peer ports use '
                    'native fibre channel protocol, not FICON, you can still put all the peer-to-peer ports in one '
                    'domain index zone. Consider creating a separate logical switch for disk mirroring ports.'),
            dict(p1=''),
            dict(p1='Any channel or control unit interface that was offline when the data was collected will need to '
                    'be added.'),
            dict(p1=''),
            dict(p1='This is a little more flexible than "By_Channel" because supplying IOCPs is not required and a '
                    'new IODF can be activated with new link addresses for existing CHPIDs and control units. If there '
                    'are new CHPIDs, new devices, or any channel or device connection was moved, the zone database '
                    'will need to be updated accordingly.'),
        ],
        a=_zone_option_dev_func,
        sname='By_Device'
    ),
    dict(
        h=[
            dict(p1='Zones are built on a per channel basis. One peer-zone per CHPID is created. The member is the '
                    'CHPID port. The principal members are all the ports associated with the "LINK=" statements in the '
                    'CNTLUNIT macros. Since the ports are determined from the "LINK=" statements, the control units do '
                    'not have to be online. Entry link addresses can only be determined when the channel is online and '
                    'the CHPIDs are configured online.'),
            dict(p1=''),
            dict(p1='A copy of the IOCP in plain text, build I/O configuration statements from HCD, from each CPC is '
                    'required as input for the -iocp command line option to build a worksheet with these zones.'),
            dict(p1=''),
            dict(p1='This method produces the smallest zones, but a zone configuration update will be required if:'),
            dict(p1='', b1_char='*', b1_prefix='  '),
            dict(b1='A new IODF is activated'),
            dict(b1='A previously offline CHPID is configured online'),
        ],
        a=_zone_option_channel_func,
        sname='By_Channel'
    ),
)

_eh_l = [
    dict(p1='**Overview**'),
    dict(p1=''),
    dict(p1='The purpose of this script is to provide a simple method to create zones for mainframe environments when '
            'SANnav is not used or when alternative zoning methods are desired. The output is a configuration workbook '
            'that can be used as input to zone_config.py.'),
    dict(p1=''),
    dict(p1='Three step approach:'),
    dict(p1=''),
    dict(b1_char='1'),
    dict(b1='Collect data using multi_capture.py or capture.py.'),
    dict(b1='Use this script to create a zone configuration workbook.'),
    dict(b1='Use zone_config.py to implement the zone configuration.'),
    dict(p1=''),
    dict(p1='Some of the zoning method options are based on connections to the fabric. Since its often not possible to '
            'have a fully configured fabric with all intended connections online, creating a zone configuration '
            'workbook allows you to modify the zone definitions to account for channels and devices not connected at '
            'the time data from switches was collected.'),
    dict(p1=''),
    dict(p1='Since the zone_config.py script has an option to output CLI, this also makes it easy for service '
            'organizations to create CLI zone scripts for SAN implementation engineers that do not have access to a '
            'Python interpreter with the brcddb and brcdapi libraries installed.'),
    dict(p1=''),
    dict(p1='Zone configurations are based on a fabric. Fabrics are identified by a world wide name (WWN), -wwn '
            'parameter. Use the -scan parameter to determine the WWN.'),
    dict(p1=''),
    dict(p1='**Naming Conventions**'),
    dict(p1='', b1_char='*', b1_prefix='  ', b2_char='-', b2_prefix='      '),
    dict(b1='A unique zone configuration name, -cfg, is enforced.'),
    dict(b1='Where applicable, zone name and alias names are a hash of the serial number, CHPID tag, and device type.'),
    dict(b1='Creating unique zone database items make it easy to:'),
    dict(b2='Back out zoning changes.'),
    dict(b2='Clean out previous zone databases by simply running the zone_config.py script with "full_purge"'),
    dict(p1=''),
    dict(p1='**Output File**'),
    dict(p1=''),
    dict(p1='There is a sheet for each zoning method. If IOCPs were not provided, the "By_Channel" work sheet will be '
            'empty. See "Zoning Methods" for details.'),
    dict(p1=''),
    dict(p1='When creating a zone object, If the zone object name already exists, a new one is created with a unique '
            'number appended to ensure it has a unique name. This makes it easy to use the "full_purge" action in '
            'zone_config.py workbooks to clean out no longer used zone database items after the zone configuration is '
            'validated.'),
    dict(p1=''),
    dict(p1='**Zoning Tips**'),
    dict(p1=''),
    dict(p1='Create a new zone configuration whenever modifying the zone database of an existing fabric so that your '
            'back-out plan is to simply activate the old zone configuration.'),
    dict(p1=''),
    dict(p1='**What Is Zoning?**'),
    dict(p1=''),
    dict(p1='SAN switches do not route frames anywhere until programmed to do so. This is referred to as "Zoning". '
            'In mainframe environments, controlling what servers (channels) can access what devices (control units) is '
            'handled by the IODF. In open systems environments, this function is handled by the SAN. As a result, there'
            ' are two types of zoning, WWN and Domain Index.'),
    dict(p1=''),
    dict(p1='Some of the zoning methods use peer zones. There is no need for mainframe storage administrators to know '
            'the difference between a zone and a peer zone. For the curious:'),
    dict(p1='', b1_prefix='  ', b1_char='*'),
    dict(b1='Peer zones have two membership lists, member and principal member.'),
    dict(b1='Any port in the member list can communicate with any port in the principal list and vice versa.'),
    dict(b1='Ports in the member list cannot communicate with other ports in the member list.'),
    dict(b1='Ports in the principal member list cannot communicate with other ports in the principal member list.'),
    dict(p1=''),
    dict(p1='The two types of zoning are:'),
    dict(p1=''),
    dict(p1='WWN'),
    dict(p1=''),
    dict(p1='Used by open systems. All Fibre Channel devices have a unique name. This name is referred to as a World '
            'Wide Name (WWN). In open systems environments, zones consist of WWNs. Devices can only communicate with '
            'other devices whose WWNs are in the same zone.'),
    dict(p1=''),
    dict(p1='Domain Index (d,i)'),
    dict(p1=''),
    dict(p1='Used by mainframes. In mainframe environments, controlling what channels can access what control units is '
            'handled by the IODF. The IODF is based on link addresses. Link addresses correspond to physical ports in '
            'the SAN. Mainframe zoning, therefore, is based on physical ports. Brocade offers flexible port numbering '
            'options. To ensure a consistent way to reference physical ports, Brocade SAN switches assign a consistent '
            'port index to each physical port. There is no need for a mainframe SAN administrator to know anything else'
            ' about port numbering options or port indexes. It is mentioned here only to explain why the index is used '
            'instead of the port number.'),
    dict(p1=''),
    dict(p1='Since FICON is a protocol transported by Fibre Channel, channels and control unit ports do have WWNs, but '
            'they are not obvious in mainframe tools and, more importantly, the IODF is built by link address, not '
            'WWN.'),
    dict(p1=''),
    dict(p1='The most common approach to zoning for mainframes has been to create a single large zone with all ports in'
            ' it. Since all ports are in a zone, all link addresses are accessible. Large zones have become somewhat of'
            ' a concern. The purpose of this script is to facilitate different zoning methods, some of which are '
            'intended to reduce the size of mainframe zones. SAN administrators should discuss these zoning options '
            'with their SAN vendor or Brocade representative.'),
    dict(p1=''),
    dict(p1='**Zoning Methods**'),
]
# Add the Notes page, including the zone options. While in this loop, we'll add the zone options to the extended help
_general_notes_l = [
    [dict(t='Tab', font=_bold_font), dict(t='Description', font=_bold_font)]
]
for _d in _zone_options_l:
    _eh_l.extend([dict(p1=''), dict(p1='**' + _d['sname'] + '**'), dict(p1=''),])
    _buf_l = gen_util.convert_to_list(_d['h'])
    _eh_l.extend(_buf_l)
    _buf_l = gen_util.format_text(_buf_l)
    _general_notes_l.append([dict(t=_d['sname']), dict(t=None if len(_buf_l) == 0 else _buf_l[0])])
    if len(_buf_l) > 1:
        for _buf in _buf_l[1:]:
            _general_notes_l.append([dict(), dict(t=_buf)])
_eh_l.append(dict(p1=''))

# Input parameter definitions
_input_d = dict(
    i=dict(
        r=False,
        h='Required unless using -eh. Name of input file generated by capture.py, combine.py, or multi_capture.py. '
          'Extension ".json" is automatically added.'
    ),
    o=dict(
        r=False,
        h='Required unless using -eh or -scan. Output file file name. Extension ".xlsx" is automatically added.'
    ),
    wwn=dict(r=False, h='Required unless using -eh or -scan. Fabric WWN. Run with -scan option if unknown.'),
    cfg=dict(
        r=False,
        h='Required unless using -eh or -scan. This name of the new zone configuration to create.'
    ),
    iocp=dict(
        r=False,
        h='Optional. Name of folder with IOCP files. If this option is specified, a worksheet with zone configurations '
          'based on channel paths is created. All files in this folder must be IOCP files in plain text (build I/O '
          'configuration statements from HCD) and must begin with the CPC serial number followed by "_". Leading 0s are'
          ' not required. The remaining characters in the file name are not used by this script. Example: for a CPC '
          'with serial number 12345: 12345_M90_iocp.txt.'
    )
)
_input_d.update(gen_util.parseargs_scan_d.copy())
_input_d.update(gen_util.parseargs_eh_d.copy())
_input_d.update(gen_util.parseargs_log_d.copy())


def pseudo_main(fabric_obj, output_file, zonecfg, iocp_l):
    """Basically the main(). Did it this way so that it can easily be used as a standalone module or called externally.

    :param fabric_obj: Fabric object
    :type fabric_obj: brcddb.classes.fabric.FabricObj
    :param output_file: Name of zone configuration workbook
    :type output_file: str
    :param zonecfg: Zone configuration name
    :type zonecfg: str
    :param iocp_l: IOCP files to read
    :type iocp_l: list
    :return: Exit code. See exist codes in brcddb.brcddb_common
    :rtype: int
    """
    global _zone_options_l, _general_notes_l

    ec, notes_l = brcddb_common.EXIT_STATUS_OK, list()

    # Build cross-references and custom search terms.
    brcdapi_log.log('Building cross-references', echo=True)
    proj_obj = fabric_obj.r_project_obj()
    brcddb_project.build_xref(proj_obj)
    brcddb_project.add_custom_search_terms(proj_obj)
    for file in iocp_l:
        brcddb_iocp.parse_iocp(proj_obj, file)
    rnid_d = brcddb_fabric.rnid_groups(fabric_obj)

    # Create a workbook and add a worksheet for each zone method
    brcdapi_log.log('Creating the zone configuration workbook, ' + output_file, echo=True)
    wb = brcddb_zone.create_zone_workbook()
    # Insert the "About" worksheet
    buf_l = (
        'The contents of this Workbook were generated from a script that reads a data from a data collection file or '
        'directly from a switch. It also reads IOCP plain text files. It generates a workbook to be used with '
        'zone_config.py. The zone and alias naming convention is a hash of the device type, serial number, and tag.'
        '',
        'See the notes page for details regarding the different worksheets and naming conventions.',
    )
    report_utils.about_page(wb, 1, 'About', os.path.basename(__file__), __version__, buf_l)
    no_zone_fab_obj = brcddb_fabric.copy_fab_obj(fabric_obj, skip_alias=True, skip_zone=True, skip_zonecfg=True)
    for zone_d in _zone_options_l:
        brcddb_zone.add_zone_worksheet(
            wb,
            zone_d['sname'],
            zone_d['a'](no_zone_fab_obj, zonecfg, rnid_d, notes_l),
            1)

    # Add the notes sheet
    sheet = wb.create_sheet(index=1, title='Notes')
    sheet.page_setup.paperSize = sheet.PAPERSIZE_LETTER
    for col in range(0, len(_zonecfg_col)):
        sheet.column_dimensions[xl.get_column_letter(col+1)].width = _zonecfg_col[col]

    # Add the content to the Notes sheet
    row = 1
    for buf in notes_l:
        excel_util.cell_update(sheet, row, 1, buf, font=_std_font, align=_align_wrap)
        sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=2)
        row += 1
    if len(notes_l) > 0:
        row += 1
    for row_l in _general_notes_l:
        col = 1
        for col_d in row_l:
            excel_util.cell_update(
                sheet,
                row,
                col,
                col_d.get('t'),
                font=col_d.get('font', _std_font),
                align=col_d.get('align', _align_wrap),
                border=col_d.get('border'),
                link=col_d.get('link')
            )
            col_span = col_d.get('span', 1)
            if col_span > 1:
                sheet.merge_cells(start_row=row, start_column=col, end_row=row, end_column=col + col_span-1)
            col += col_span
        row += 1

    # Save the workbook
    if not excel_util.save_report(wb, output_file, 'Consoli-Solutions'):
        ec = brcddb_common.EXIT_STATUS_INPUT_ERROR

    return ec


def _get_input():
    """Parses the module load command line

    :return: Exit code. See exist codes in brcddb.brcddb_common
    :rtype: int
    """
    global __version__, _input_d, _eh_l
    
    ec = brcddb_common.EXIT_STATUS_OK
    error_msg = ''

    # Get command line input
    buf = 'Creates workbooks for use with zone_config.py for common mainframe zoning operations. Each workbook contains'
    buf += ' a worksheet for each type of zone configuration. SAN administrators pick the worksheet with their '
    buf += 'preferred zoning method.'
    args_d = gen_util.get_input(buf, _input_d)

    # Set up logging
    brcdapi_log.open_log(
        folder=args_d['log'],
        suppress=args_d['sup'],
        no_log=args_d['nl'],
        version_d=brcdapi_util.get_import_modules()
    )

    # Setup feedback messages for options
    feedback_d = dict(
        i=dict(h=str(args_d['i']), rl=('eh',), e='ERROR: Missing required input parameter.'),
        o=dict(h=str(args_d['o']), rl=('eh', 'scan'), e='ERROR: Missing required input parameter.'),
        wwn=dict(h=str(args_d['wwn']), rl=('eh', 'scan'), e='ERROR: Missing required input parameter.'),
        cfg=dict(h=str(args_d['cfg']), rl=('eh', 'scan'), e='ERROR: Missing required input parameter.'),
    )
    for key, d in feedback_d.items():
        try:
            for required_key in d['rl']:
                if args_d[required_key]:
                    raise Found  # The parameter is not required: -eh or -scan
            if args_d[key] is None:
                d['h'] += ' ' + d['e']
                ec = brcddb_common.EXIT_STATUS_INPUT_ERROR
        except Found:
            continue

    # Read and validate the project file, -i
    # proj_obj gets over-written. It needs to be initialized with something.
    proj_obj = fos_auth.create_error(brcdapi_util.HTTP_BAD_REQUEST, 'Unknown', None)
    fabric_obj = None
    if args_d['i'] is not None and not args_d['eh']:
        in_file = brcdapi_file.full_file_name(args_d['i'], '.json')
        proj_obj = brcddb_project.read_from(in_file)
        if fos_auth.is_error(proj_obj):
            error_msg = fos_auth.formatted_error_msg(proj_obj)
            feedback_d['i']['h'] += ' ERROR: reading file. See below for details.'
            ec = brcddb_common.EXIT_STATUS_INPUT_ERROR

        # Validate the -wwn parameter
        elif args_d['wwn'] is not None:
            fabric_obj = proj_obj.r_fabric_obj(args_d['wwn'])
            if fabric_obj is None:
                feedback_d['wwn']['h'] += ' ERROR: Does not exist in ' + args_d['i']
                ec = brcddb_common.EXIT_STATUS_INPUT_ERROR

    # Read the directory of IOCP files, -iocp
    iocp_l = list()
    if args_d['iocp'] is not None:
        iocp_l = brcdapi_file.read_directory(args_d['iocp'], full=True)
        if len(iocp_l) == 0:
            feedback_d['iocp']['h'] += ' ERROR: Folder not found or folder is empty.'
            ec = brcddb_common.EXIT_STATUS_INPUT_ERROR

    # Make sure the zone configuration name is unique
    if fabric_obj is not None and isinstance(args_d['cfg'], str):
        zonecfg = brcddb_fabric.unique_zonecfg_name(fabric_obj, args_d['cfg'])
        if zonecfg != args_d['cfg']:
            feedback_d['cfg']['h'] = ' ERROR:' + args_d['cfg'] + ' already used. Suggestion: ' + zonecfg
            ec = brcddb_common.EXIT_STATUS_INPUT_ERROR

    # Command line feedback
    ml = [
        '',
        os.path.basename(__file__) + ', ' + __version__,
        'Input file, -i:           ' + feedback_d['i']['h'],
        'Output file, -o:          ' + feedback_d['o']['h'],
        'Fabric WWN, -wwn:         ' + feedback_d['wwn']['h'],
        'Zone configuration, -cfg: ' + feedback_d['cfg']['h'],
        'IOCP, -iocp:              ' + str(args_d['iocp']),
        'Extended help, -eh:       ' + str(args_d['eh']),
        'Scan, -scan:              ' + str(args_d['scan']),
        'Log, -log:                ' + str(args_d['log']),
        'No log, -nl:              ' + str(args_d['nl']),
        'Suppress, -sup:           ' + str(args_d['sup']),
        '',
        error_msg,
        ]
    brcdapi_log.log(ml, echo=True)

    # Extended help, Scan, and error bail out.
    if args_d['eh']:
        brcdapi_log.log(gen_util.format_text(_eh_l), echo=True)
        ec = brcddb_common.EXIT_STATUS_INPUT_ERROR
    if args_d['scan'] and not fos_auth.is_error(proj_obj):
        brcdapi_log.log(brcddb_project.scan(proj_obj, fab_only=False, logical_switch=True), echo=True)
        ec = brcddb_common.EXIT_STATUS_INPUT_ERROR
    if not isinstance(args_d['i'], str) or not isinstance(args_d['o'], str):
        ec = brcddb_common.EXIT_STATUS_INPUT_ERROR

    return ec if ec != brcddb_common.EXIT_STATUS_OK else \
        pseudo_main(fabric_obj, brcdapi_file.full_file_name(args_d['o'], '.xlsx'), args_d['cfg'], iocp_l)


##################################################################
#
#                    Main Entry Point
#
###################################################################
if _DOC_STRING:
    print('_DOC_STRING is True. No processing')
    exit(0)

if _STAND_ALONE:
    _ec = _get_input()
    brcdapi_log.close_log(['', 'Processing Complete. Exit code: ' + str(_ec)], echo=True)
    exit(_ec)
